import xml.etree.ElementTree as ET

import pandas as pd
import requests
from openpyxl.styles import PatternFill

from config import (
    PDF_ROOT,
    XML_ROOT,
    VERIFICATION_ROOT,
    YEAR_INTS,
    ensure_folders,
)


HEADERS = {"User-Agent": "Mozilla/5.0"}


def run() -> pd.DataFrame:
    ensure_folders()
    summary = []

    for year in YEAR_INTS:
        print(f"\n===== STAGE 2: {year} =====")

        xml_url = (
            "https://disclosures-clerk.house.gov/"
            f"public_disc/financial-pdfs/{year}FD.xml"
        )

        response = requests.get(
            xml_url,
            headers=HEADERS,
            timeout=60,
        )
        response.raise_for_status()

        xml_path = XML_ROOT / f"{year}FD.xml"
        xml_path.write_bytes(response.content)

        root = ET.parse(xml_path).getroot()

        all_rows = []
        for member in root.iter("Member"):
            all_rows.append({
                child.tag: child.text
                for child in member
            })

        full_index_df = pd.DataFrame(all_rows)

        if "DocID" in full_index_df.columns:
            full_index_df["DocID"] = (
                full_index_df["DocID"]
                .fillna("")
                .astype(str)
                .str.strip()
            )

        ptr_mask = (
            full_index_df["FilingType"]
            .fillna("")
            .astype(str)
            .str.strip()
            .eq("P")
        )

        ptr_df = full_index_df.loc[ptr_mask].copy()

        year_folder = PDF_ROOT / str(year)
        year_folder.mkdir(parents=True, exist_ok=True)

        pdf_docids = {
            pdf.stem
            for pdf in year_folder.glob("*.pdf")
        }

        full_index_df["PDF in Drive?"] = ""

        for row_number in full_index_df.index:
            filing_type = str(
                full_index_df.loc[row_number, "FilingType"]
            ).strip()
            doc_id = str(
                full_index_df.loc[row_number, "DocID"]
            ).strip()

            if filing_type == "P":
                full_index_df.loc[
                    row_number,
                    "PDF in Drive?"
                ] = "YES" if doc_id in pdf_docids else "NO"

        expected_docids = set(ptr_df["DocID"].astype(str))
        missing_docids = sorted(expected_docids - pdf_docids)
        extra_docids = sorted(pdf_docids - expected_docids)

        expected_count = len(expected_docids)
        pdf_count = len(pdf_docids)
        matched_count = len(expected_docids & pdf_docids)
        missing_count = len(missing_docids)
        extra_count = len(extra_docids)

        complete = missing_count == 0 and extra_count == 0

        verification_df = pd.DataFrame({
            "Check": [
                "Year",
                "PTRs in official House XML",
                "PDF files in Drive",
                "Matched PTR PDFs",
                "Missing PTR PDFs",
                "Extra PDFs not in XML",
                "Complete?",
            ],
            "Result": [
                year,
                expected_count,
                pdf_count,
                matched_count,
                missing_count,
                extra_count,
                "YES" if complete else "NO",
            ],
        })

        excel_path = VERIFICATION_ROOT / f"House_PTR_Verification_{year}.xlsx"

        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            full_index_df.to_excel(
                writer,
                sheet_name="House XML Index",
                index=False,
            )
            verification_df.to_excel(
                writer,
                sheet_name="Verification",
                index=False,
            )

            if missing_count:
                pd.DataFrame({
                    "Missing DocID": missing_docids
                }).to_excel(
                    writer,
                    sheet_name="Missing PTRs",
                    index=False,
                )

            if extra_count:
                pd.DataFrame({
                    "Extra PDF DocID": extra_docids
                }).to_excel(
                    writer,
                    sheet_name="Extra PDFs",
                    index=False,
                )

            worksheet = writer.book["House XML Index"]

            headers_in_sheet = {
                cell.value: cell.column
                for cell in worksheet[1]
            }

            filing_col = headers_in_sheet["FilingType"]
            drive_col = headers_in_sheet["PDF in Drive?"]

            green_fill = PatternFill(
                fill_type="solid",
                fgColor="C6EFCE",
            )
            red_fill = PatternFill(
                fill_type="solid",
                fgColor="FFC7CE",
            )

            for excel_row in range(2, worksheet.max_row + 1):
                filing_type = worksheet.cell(
                    excel_row,
                    filing_col,
                ).value
                drive_status = worksheet.cell(
                    excel_row,
                    drive_col,
                ).value

                if filing_type == "P":
                    fill = (
                        green_fill
                        if drive_status == "YES"
                        else red_fill
                    )
                    for cell in worksheet[excel_row]:
                        cell.fill = fill

        summary.append({
            "Year": year,
            "House PTRs": expected_count,
            "PDFs in Drive": pdf_count,
            "Matched": matched_count,
            "Missing": missing_count,
            "Extra": extra_count,
            "Complete": complete,
        })

        print("House PTRs:", expected_count)
        print("PDFs in Drive:", pdf_count)
        print("Matched:", matched_count)
        print("Missing:", missing_count)
        print("Extra:", extra_count)
        print("Complete:", complete)
        print("Excel:", excel_path)

    summary_df = pd.DataFrame(summary)
    print("\n===== ALL YEARS =====")
    print(summary_df.to_string(index=False))
    return summary_df


if __name__ == "__main__":
    run()
